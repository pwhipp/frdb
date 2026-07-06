from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


CANONICAL_WORKBOOK_RELATIVE_PATH = Path("source_materials/interactive_evidence_recovery_comparison.xlsx")
DEFAULT_COMPARISONS_SHEET = "Comparisons"


@dataclass(frozen=True)
class CombinedWorkbook:
    path: Path
    study_sheet: str
    comparisons_sheet: str
    sheet_names: tuple[str, ...]


def default_workbook_path(repo_root: Path) -> Path:
    return repo_root / CANONICAL_WORKBOOK_RELATIVE_PATH


def resolve_workbook_path(repo_root: Path, workbook_path: Path | None) -> Path:
    return (workbook_path or default_workbook_path(repo_root)).resolve()


def inspect_combined_workbook(path: Path, comparisons_sheet: str) -> CombinedWorkbook:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as error:
        raise ValueError(f"Could not open workbook {path}: {error}") from error

    try:
        sheet_names = tuple(workbook.sheetnames)
    finally:
        workbook.close()

    if not sheet_names:
        raise ValueError(f"{path} has no worksheets")
    if comparisons_sheet not in sheet_names:
        raise ValueError(
            f"{path} has no {comparisons_sheet!r} worksheet; available worksheets: {list(sheet_names)}"
        )
    if sheet_names[0] == comparisons_sheet:
        raise ValueError(
            f"{path} first worksheet must contain study rows because the parser reads the first worksheet as studies"
        )

    return CombinedWorkbook(
        path=path,
        study_sheet=sheet_names[0],
        comparisons_sheet=comparisons_sheet,
        sheet_names=sheet_names,
    )
